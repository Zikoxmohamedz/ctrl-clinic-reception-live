import json,re,unicodedata
from pathlib import Path
from collections import defaultdict
from difflib import SequenceMatcher
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
root=Path('.local');data=json.loads((root/'accounting-source.json').read_text(encoding='utf-8'))
review=json.loads((root/'negative-review-data.json').read_text(encoding='utf-8'))
def sheet(wb,name,headers,rows):
 ws=wb.create_sheet(name);ws.sheet_view.rightToLeft=True;ws.append(headers)
 for row in rows:ws.append(row)
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='183153')
 for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=42 if 'اسم' in h or 'ملاحظ' in h or 'صنف' in h or 'سبب' in h else 23
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if isinstance(cell.value,(float,int)):cell.number_format='0.####'
   elif isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')):cell.data_type='s'
 return ws
wb=Workbook();wb.remove(wb.active)
sheet(wb,'طريقة المراجعة',['البند','التوضيح'],[
 ['الفترة','31/7/2026 حتى 27/9/2026؛ بعد إقفال مدينة نصر وتصحيح افتتاح حلوان'],
 ['المعادلة','الافتتاح + الوارد − الصرف − التحويل + فروق الجرد = الرصيد'],
 ['السالب','سالب دفتري يحتاج مراجعة مستندات ووحدة القياس؛ ليس إثباتًا لعجز فعلي'],
 ['الحركات','تفاصيل الحركات للأصناف السالبة فقط. استبعاد حركة يعني أنها سبقت تأسيس الرصيد'],
 ['التصحيح','راجع هل الجرد بالعبوة والصرف بالوحدة، وهل يوجد وارد ناقص أو صرف مكرر؛ لا تضرب الكميات تلقائيًا']])
keys=['branch','code','name','unit','opening','incoming','outgoing','transfers','count_adjustment','balance','last_count','last_count_date']
sheet(wb,'الأرصدة السالبة',['الفرع','الكود','اسم الصنف','الوحدة','افتتاح','وارد','صرف','تحويل','فروق الجرد','الرصيد السالب','آخر كمية فعلية','تاريخ آخر جرد'],[[r.get(k) for k in keys] for r in review['negatives']])
keys=['branch','code','name','unit','date','type','quantity','before','after','excluded','reference','notes'];types={'addition':'وارد','consumption':'صرف','transfer':'تحويل صادر','opening':'افتتاح','count':'جرد'}
sheet(wb,'حركات الأصناف السالبة',['الفرع','الكود','اسم الصنف','الوحدة','التاريخ','نوع الحركة','الكمية','قبل','بعد','مستبعدة قبل التأسيس','مرجع الحركة','ملاحظات'],[[types.get(r.get(k),r.get(k)) if k=='type' else ('نعم' if r.get(k) else 'لا') if k=='excluded' else r.get(k) for k in keys] for r in review['movements']])
negative_path=root/'مراجعة-الأرصدة-السالبة-27-9-2026.xlsx';wb.save(negative_path)
materials=[m for m in data['materials'] if not m.get('archived_at')]
def normalize(s):return re.sub(r'[\W_]+','',unicodedata.normalize('NFKC',str(s or '')).casefold())
issues=[];exact=defaultdict(list);names=defaultdict(list);codes=defaultdict(list)
for m in materials:
 exact[(normalize(m['name']),normalize(m['unit']))].append(m);names[normalize(m['name'])].append(m);codes[str(m['code']).strip().casefold()].append(m)
def issue(kind,a,b=None):issues.append([kind,a['code'],a['name'],a['unit'],b['code'] if b else '',b['name'] if b else '',b['unit'] if b else '','للمراجعة فقط؛ لم يتم الدمج أو تغيير الوحدات'])
for group in exact.values():
 for m in group[1:]:issue('اسم ووحدة متطابقان بعد توحيد الكتابة',group[0],m)
for group in names.values():
 for m in group[1:]:
  if normalize(group[0]['unit'])!=normalize(m['unit']):issue('اسم واحد بوحدات مختلفة',group[0],m)
for group in codes.values():
 for m in group[1:]:issue('كود مكرر باختلاف الكتابة',group[0],m)
for i,a in enumerate(materials):
 for b in materials[i+1:]:
  na,nb=normalize(a['name']),normalize(b['name'])
  if na!=nb and normalize(a['unit'])==normalize(b['unit']) and re.findall(r'\d+(?:\.\d+)?',a['name'])==re.findall(r'\d+(?:\.\d+)?',b['name']) and SequenceMatcher(None,na,nb).ratio()>=.94:issue('اسمان متشابهان جدًا؛ تحقق من المنتج',a,b)
wb=Workbook();wb.remove(wb.active)
sheet(wb,'ملاحظات',['البند','النتيجة'],[['أصناف نشطة',len(materials)],['اشتباه تكرار أو تشابه',len(issues)],['أصناف مؤقتة',sum(bool(m.get('is_temp')) for m in materials)],['تكلفة صفر أو غير مسجلة',sum(not m.get('cost_price') for m in materials)],['حدود الفحص','تشابه الاسم لا يثبت أنه نفس المنتج؛ قارن التركيز والحجم والشركة والوحدة قبل الدمج']])
sheet(wb,'اشتباه التكرار',['سبب المراجعة','الكود الأول','اسم الصنف الأول','وحدته','الكود الثاني','اسم الصنف الثاني','وحدته','ملاحظات'],issues)
sheet(wb,'دليل الأصناف',['الكود','اسم الصنف','الوحدة','الفئة','تكلفة الوحدة الحالية','سعر البيع الافتراضي المستقل','مؤقت'],[[m['code'],m['name'],m['unit'],m.get('category'),m.get('cost_price'),m.get('default_price'),'نعم' if m.get('is_temp') else 'لا'] for m in materials])
catalog_path=root/'مراجعة-دليل-الأصناف-27-9-2026.xlsx';wb.save(catalog_path)
wb=Workbook();wb.remove(wb.active);sheet(wb,'تسعير الخامات',['code','name','unit','category','cost_price'],[[m['code'],m['name'],m['unit'],m.get('category'),m.get('cost_price') or None] for m in materials]);wb.save(root/'خامات-جاهزة-للتسعير.xlsx')
assert load_workbook(negative_path).worksheets[1].max_row==len(review['negatives'])+1
(root/'catalog-review-summary.json').write_text(json.dumps({'materials':len(materials),'issues':issues,'zeroCost':sum(not m.get('cost_price') for m in materials)},ensure_ascii=False),encoding='utf-8')
print(json.dumps({'negativeRows':len(review['negatives']),'catalogIssues':len(issues),'activeMaterials':len(materials),'zeroCost':sum(not m.get('cost_price') for m in materials)}))
